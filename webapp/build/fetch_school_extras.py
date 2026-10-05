#!/usr/bin/env python3
"""NYC Daycare Check — school-context extras for school-based Pre-K pages.

Three DOE InfoHub workbooks, all current where the Open Data mirrors froze
years ago (survey mirror stopped at 2019, demographics at 2021-22):

    guardian survey  2026-public-data-file-guardian.xlsx      (SY 2025-26)
    demographics     demographic-snapshot-2021-22-to-2025-26-public.xlsx
    attendance       public-school-attendance-results-2019-2025.xlsx

Parsed into data/processed/school_extras.json keyed by DBN, consumed by
build_data.py for prek-tier records whose dbn is a real school DBN
(##B###). Coverage measured 2026-09-07 against our 622 school DBNs:
guardian 608 (98%). The B-5 early-childhood survey file (NYCEEC sites) is
deliberately NOT ingested: it keys by C-code with no address, and exact
name matching reached only 15% of our CBO pages — below the bar for
attaching survey numbers to named businesses. Revisit if the DOE publishes
a C-code crosswalk.

Annual files: the fetch is skipped when the sidecar is younger than
REFRESH_AFTER_DAYS, so the nightly refresh doesn't re-download ~150MB of
workbooks that change once a year.
"""
import io
import json
import os
import re
import sys
import time
import urllib.request
from datetime import datetime, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.normpath(os.path.join(HERE, "..", "data", "processed", "school_extras.json"))
UA = "NYCDaycareCheck/1.0 (+https://nycdaycarecheck.com; public-data ingest)"
REFRESH_AFTER_DAYS = 30

URLS = {
    "guardian": "https://infohub.nyced.org/docs/default-source/default-document-library/2026-public-data-file-guardian.xlsx",
    "demo": "https://infohub.nyced.org/docs/default-source/default-document-library/demographic-snapshot-2021-22-to-2025-26-public.xlsx",
    "attendance": "https://saintrafileprod01.blob.core.windows.net/prd-intra/docs/default-source/large-files/public-school-attendance-results-2019-2025.xlsx",
}

DBN_RE = re.compile(r"^\d{2}[KMQXR]\d{3}$")


def _fetch(url, retries=3):
    for i in range(retries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA})
            with urllib.request.urlopen(req, timeout=120) as r:
                return r.read()
        except Exception as e:
            if i == retries - 1:
                raise
            print(f"  retry {i + 1} after {type(e).__name__}", flush=True)
            time.sleep(10)


def _raw(v):
    """Excel cell -> full-precision float or None (or an Above/Below string)."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or s.upper() in ("N/A", "NA", "S", "--"):
        return None
    if s.lower().startswith(("above", "below")):
        return s  # privacy-suppressed label, passed through verbatim
    try:
        return float(s.rstrip("%"))
    except ValueError:
        return None


def _num(v):
    n = _raw(v)
    if n is None or isinstance(n, str):
        return None
    return round(n, 1)


def _pct(v):
    """% columns arrive as fractions (0.4149) or percents (41.5 / '41.5%').
    Convert BEFORE rounding — rounding a fraction first destroys precision.
    Privacy labels ('Above 95%', 'Below 5%') pass through as strings."""
    n = _raw(v)
    if n is None:
        return None
    if isinstance(n, str):
        return n
    return round(n * 100, 1) if n <= 1.0 else round(n, 1)


def parse_guardian(blob):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True)
    ws = wb["Total"]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(c or "").strip() for c in next(rows)]

    def col(prefix):
        for i, h in enumerate(hdr):
            if h.startswith(prefix):
                return i
        return None

    c_resp = col("Total Family Response Count")
    c_rate = col("Total Family Response Rate")
    c_sat = col("Family Satisfaction with Child")
    c_trust = col("Parent-Teacher Trust")
    out = {}
    for r in rows:
        dbn = str(r[0] or "").strip().upper()
        if not DBN_RE.match(dbn):
            continue
        out[dbn] = {
            "famSat": _num(r[c_sat]) if c_sat is not None else None,
            "ptTrust": _num(r[c_trust]) if c_trust is not None else None,
            "respCount": _num(r[c_resp]) if c_resp is not None else None,
            "respRate": _pct(r[c_rate]) if c_rate is not None else None,
        }
    wb.close()
    return out


def parse_demo(blob):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True)
    ws = wb["School"]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(c or "").strip() for c in next(rows)]
    ix = {h: i for i, h in enumerate(hdr)}
    out = {}
    for r in rows:
        dbn = str(r[0] or "").strip().upper()
        if not DBN_RE.match(dbn):
            continue
        year = str(r[ix["Year"]] or "")
        prev = out.get(dbn)
        if prev and prev["year"] >= year:   # keep latest school year only
            continue
        out[dbn] = {
            "year": year,
            "enroll": _num(r[ix["Total Enrollment"]]),
            "g3k": _num(r[ix["Grade 3K"]]),
            "gpk": _num(r[ix["Grade PK (Half Day & Full Day)"]]),
            "pct": {
                "asian": _pct(r[ix["% Asian and Pacific Islander"]]),
                "black": _pct(r[ix["% Black"]]),
                "hispanic": _pct(r[ix["% Hispanic"]]),
                "white": _pct(r[ix["% White"]]),
                "multi": _pct(r[ix["% Multi-Racial"]]),
                "native": _pct(r[ix["% Native American"]]),
                "swd": _pct(r[ix["% Students with Disabilities"]]),
                "ell": _pct(r[ix["% English Language Learners"]]),
                "poverty": _pct(r[ix["% Poverty"]]),
            },
        }
    wb.close()
    return out


def parse_attendance(blob):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True)
    ws = wb["All Students"]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(c or "").strip() for c in next(rows)]
    ix = {h: i for i, h in enumerate(hdr)}
    c_att = next((i for h, i in ix.items() if h.startswith("% Attendance")), None)
    c_chr = next((i for h, i in ix.items() if "Chronically Absent" in h and h.startswith("%")), None)
    out = {}
    for r in rows:
        dbn = str(r[0] or "").strip().upper()
        if not DBN_RE.match(dbn):
            continue
        if str(r[ix.get("Grade", 2)] or "") != "All Grades":
            continue
        year = str(r[ix["Year"]] or "")
        prev = out.get(dbn)
        if prev and prev["year"] >= year:
            continue
        out[dbn] = {
            "year": year,
            "rate": _num(r[c_att]) if c_att is not None else None,
            "chronic": _num(r[c_chr]) if c_chr is not None else None,
        }
    wb.close()
    return out


def main():
    force = "--force" in sys.argv
    if not force and os.path.exists(OUT):
        try:
            cur = json.load(open(OUT, encoding="utf-8"))
            age = (datetime.now(timezone.utc)
                   - datetime.fromisoformat(cur["fetched"])).days
            if age < REFRESH_AFTER_DAYS:
                print(f"school extras: sidecar is {age}d old (<{REFRESH_AFTER_DAYS}d) — skipping fetch")
                return
        except Exception:
            pass

    print("school extras: downloading 3 InfoHub workbooks…", flush=True)
    guardian = parse_guardian(_fetch(URLS["guardian"]))
    print(f"  survey: {len(guardian)} schools", flush=True)
    demo = parse_demo(_fetch(URLS["demo"]))
    print(f"  demographics: {len(demo)} schools", flush=True)
    att = parse_attendance(_fetch(URLS["attendance"]))
    print(f"  attendance: {len(att)} schools", flush=True)

    schools = {}
    for dbn in set(guardian) | set(demo) | set(att):
        schools[dbn] = {
            "survey": guardian.get(dbn),
            "demo": demo.get(dbn),
            "att": att.get(dbn),
        }
    payload = {
        "fetched": datetime.now(timezone.utc).isoformat(),
        "surveyYear": "2025-26",
        "schools": schools,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(payload, open(OUT, "w", encoding="utf-8"))
    print(f"school extras: wrote {len(schools)} schools -> {OUT}")


if __name__ == "__main__":
    main()
